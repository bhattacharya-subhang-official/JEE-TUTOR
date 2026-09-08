"""Upload handling — save file, extract text (pdf/docx/txt/csv), keep images for vision."""
import io
import json
import uuid
import pathlib
from config import MEDIA_DIR

UPLOAD_DIR = MEDIA_DIR / "uploads"
ALLOWED_IMAGE = {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/gif", "image/heic", "image/heif"}
MAX_DOC_CHARS = 24000


def save_and_parse(filename: str, mime: str, data: bytes) -> dict:
    uid = uuid.uuid4().hex[:10]
    safe = pathlib.Path(filename or "file").name
    ext = pathlib.Path(safe).suffix.lower()
    is_image = (mime in ALLOWED_IMAGE) or (ext in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".heic"})

    if is_image:
        file_name = f"{uid}{ext or '.png'}"
        (UPLOAD_DIR / file_name).write_bytes(data)
        meta = {"id": uid, "name": safe, "kind": "image", "mime": mime or "image/png", "file": file_name, "size": len(data)}
        (UPLOAD_DIR / f"{uid}.meta.json").write_text(json.dumps(meta))
        return {"id": uid, "name": safe, "kind": "image", "mime": meta["mime"], "size": len(data),
                "url": f"/media/uploads/{file_name}", "text": None}

    # docs → extract text
    text = ""
    try:
        if mime == "application/pdf" or ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            pages = []
            for i, page in enumerate(reader.pages[:30]):
                pages.append(f"[page {i+1}]\n{page.extract_text() or ''}")
            text = "\n".join(pages)
        elif (mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
              or ext == ".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="replace")
    except Exception as e:
        text = f"[Could not extract text: {e}]"

    file_name = f"{uid}.txt"
    (UPLOAD_DIR / file_name).write_text(text, encoding="utf-8")
    truncated = len(text) > MAX_DOC_CHARS
    meta = {"id": uid, "name": safe, "kind": "doc", "mime": mime or "text/plain", "file": file_name, "size": len(data)}
    (UPLOAD_DIR / f"{uid}.meta.json").write_text(json.dumps(meta))
    return {"id": uid, "name": safe, "kind": "doc", "mime": meta["mime"], "size": len(data),
            "url": f"/media/uploads/{file_name}",
            "text": text[:MAX_DOC_CHARS] + ("…[truncated]" if truncated else ""),
            "chars": len(text)}


def get_attachment(att_id: str):
    """Re-load an attachment by id: images → bytes for Gemini vision; docs → text."""
    meta_p = UPLOAD_DIR / f"{att_id}.meta.json"
    if not meta_p.exists():
        return None
    m = json.loads(meta_p.read_text())
    f = UPLOAD_DIR / m["file"]
    if not f.exists():
        return None
    if m["kind"] == "image":
        return {"id": att_id, "kind": "image", "name": m["name"], "mime": m["mime"],
                "url": f"/media/uploads/{m['file']}", "data": f.read_bytes()}
    return {"id": att_id, "kind": "doc", "name": m["name"], "mime": m.get("mime"),
            "url": f"/media/uploads/{m['file']}",
            "text": f.read_text(encoding="utf-8", errors="replace")[:MAX_DOC_CHARS]}
